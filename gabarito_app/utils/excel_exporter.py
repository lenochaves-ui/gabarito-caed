"""
ExcelExporter — Gera o arquivo .xlsx no mesmo formato do template CAEd
"""

import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── Paleta (idêntica ao template) ─────────────────────────────────────────────
TITLE_BG      = "1F4E79"
LP_HEADER_BG  = "2E75B6"
MAT_HEADER_BG = "ED7D31"
LP_Q_BG       = "9DC3E6"
MAT_Q_BG      = "F9CBAD"
MAT_VAL_BG    = "FCE4D6"
NR_BG         = "F2F2F2"
NR_FONT       = "999999"
NULA_BG       = "C00000"
NAME_FONT     = "1F4E79"
MAT_VAL_FONT  = "833C00"
WHITE         = "FFFFFF"

BLOCKS = [
    (1,  13, 'LP',  'Bloco 01 — Língua Portuguesa'),
    (14, 26, 'MAT', 'Bloco 02 — Matemática'),
    (27, 39, 'LP',  'Bloco 03 — Língua Portuguesa'),
    (40, 52, 'MAT', 'Bloco 04 — Matemática'),
]

thin = Side(style='thin', color="CCCCCC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def get_disc(q):
    for s, e, d, _ in BLOCKS:
        if s <= q <= e:
            return d
    return 'LP'


class ExcelExporter:

    def export(self, result_data: dict, extra_students: list = None) -> str:
        """
        Gera o .xlsx e salva no diretório padrão.
        Retorna o caminho do arquivo salvo.

        result_data: dict com 'student_name' e 'answers'
        extra_students: lista adicional de dicts no mesmo formato (para lote)
        """
        save_dir = self._get_save_dir()
        os.makedirs(save_dir, exist_ok=True)

        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        nome_curto = result_data.get('student_name', 'aluno').split()[0].upper()
        filename = f'Gabarito_{nome_curto}_{timestamp}.xlsx'
        filepath = os.path.join(save_dir, filename)

        students = [result_data]
        if extra_students:
            students.extend(extra_students)

        wb = self._build_workbook(students)
        wb.save(filepath)
        return filepath

    def _build_workbook(self, students: list) -> Workbook:
        wb = Workbook()
        ws = wb.active
        ws.title = 'Gabarito'

        # ── Linha 1: Título ────────────────────────────────────────────────────
        ws.merge_cells('A1:BA1')
        c = ws['A1']
        c.value = 'CARTÃO DE RESPOSTAS — Prefeitura Municipal de Croatá / Secretaria de Educação'
        c.font = Font(bold=True, color=WHITE, size=13)
        c.fill = PatternFill('solid', start_color=TITLE_BG)
        c.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 27.95

        # ── Linha 2: Cabeçalho de aluno + blocos ──────────────────────────────
        ws.merge_cells('A2:A3')
        c = ws['A2']
        c.value = 'ALUNO(A)'
        c.font = Font(bold=True, color=WHITE, size=11)
        c.fill = PatternFill('solid', start_color=TITLE_BG)
        c.alignment = Alignment(horizontal='center', vertical='center')

        block_cols = {
            'B01': ('B', 'N'),
            'B02': ('O', 'AA'),
            'B03': ('AB', 'AN'),
            'B04': ('AO', 'BA'),
        }
        block_meta = [
            ('B', 'N',  LP_HEADER_BG,  'BLOCO 01 — Língua Portuguesa'),
            ('O', 'AA', MAT_HEADER_BG, 'BLOCO 02 — Matemática'),
            ('AB','AN', LP_HEADER_BG,  'BLOCO 03 — Língua Portuguesa'),
            ('AO','BA', MAT_HEADER_BG, 'BLOCO 04 — Matemática'),
        ]
        for sc, ec, bg, label in block_meta:
            ws.merge_cells(f'{sc}2:{ec}2')
            c = ws[f'{sc}2']
            c.value = label
            c.font = Font(bold=True, color=WHITE, size=10)
            c.fill = PatternFill('solid', start_color=bg)
            c.alignment = Alignment(horizontal='center', vertical='center')

        ws.row_dimensions[2].height = 20.1

        # ── Linha 3: Headers Q01..Q52 ─────────────────────────────────────────
        for q in range(1, 53):
            col = q + 1
            disc = get_disc(q)
            bg = LP_Q_BG if disc == 'LP' else MAT_Q_BG
            c = ws.cell(row=3, column=col, value=f'Q{q:02d}')
            c.font = Font(bold=True, size=9)
            c.fill = PatternFill('solid', start_color=bg)
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = BORDER
        ws.row_dimensions[3].height = 18.0

        # ── Linhas de alunos ──────────────────────────────────────────────────
        name_bgs = ['D6E4F0', 'FDE8D8', 'E8F0E0', 'FDE8F8']
        for row_i, student in enumerate(students):
            r = 4 + row_i
            nome = student.get('student_name', '—')
            answers = student.get('answers', {})
            bg = name_bgs[row_i % len(name_bgs)]

            # Coluna nome
            nc = ws.cell(row=r, column=1, value=nome)
            nc.font = Font(bold=True, color=NAME_FONT, size=10)
            nc.fill = PatternFill('solid', start_color=bg)
            nc.alignment = Alignment(horizontal='left', vertical='center')
            nc.border = BORDER

            # Colunas de respostas
            for q in range(1, 53):
                col = q + 1
                val = answers.get(q, 'NR')
                disc = get_disc(q)
                cell = ws.cell(row=r, column=col)
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = BORDER

                if val == 'NR':
                    cell.value = 'N'
                    cell.font = Font(color=NR_FONT, size=10)
                    cell.fill = PatternFill('solid', start_color=NR_BG)
                elif val == 'NULA':
                    cell.value = 'NULA'
                    cell.font = Font(bold=True, color=WHITE, size=9)
                    cell.fill = PatternFill('solid', start_color=NULA_BG)
                elif val == '?':
                    cell.value = '?'
                    cell.font = Font(bold=True, color='333333', size=10)
                    cell.fill = PatternFill('solid', start_color='FFFF00')
                else:
                    cell.value = val
                    cell.font = Font(
                        bold=True,
                        color=MAT_VAL_FONT if disc == 'MAT' else '1F4E79',
                        size=10
                    )
                    cell.fill = PatternFill(
                        'solid',
                        start_color=MAT_VAL_BG if disc == 'MAT' else NR_BG
                    )

            ws.row_dimensions[r].height = 21.95

        # ── Legenda ───────────────────────────────────────────────────────────
        leg_row = 4 + len(students) + 1
        ws.cell(row=leg_row, column=1, value='LEGENDA').font = Font(
            bold=True, color=WHITE, size=10)
        ws.cell(row=leg_row, column=1).fill = PatternFill('solid', start_color=TITLE_BG)

        legend_items = [
            ('B', 'F',  'A = 1ª alternativa'),
            ('G', 'K',  'B = 2ª alternativa'),
            ('L', 'P',  'C = 3ª alternativa'),
            ('Q', 'U',  'D = 4ª alternativa'),
            ('V', 'Z',  'N = Não marcada'),
        ]
        for sc, ec, text in legend_items:
            ws.merge_cells(f'{sc}{leg_row}:{ec}{leg_row}')
            c = ws[f'{sc}{leg_row}']
            c.value = text
            c.font = Font(size=9)
            c.fill = PatternFill('solid', start_color='EAF2FB')
            c.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[leg_row].height = 20.1

        # ── Larguras de colunas ────────────────────────────────────────────────
        ws.column_dimensions['A'].width = 47.5
        for q in range(1, 53):
            ws.column_dimensions[get_column_letter(q + 1)].width = 5.4

        return wb

    def _get_save_dir(self):
        try:
            from android.storage import primary_external_storage_path
            return os.path.join(primary_external_storage_path(), 'GabaritoCAEd')
        except Exception:
            return os.path.join(os.path.expanduser('~'), 'GabaritoCAEd')
